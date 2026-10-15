import html
import hashlib
import io
import re
import string
from uuid import uuid4
import pandas as pd
import streamlit as st
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from nova.diagnostics import log_failure
from nova.services.common import BusinessError, ROLES

_MARKDOWN_SPECIAL = re.compile('([' + re.escape(string.punctuation) + '])')


def md_escape(value):
    """Show user-entered text literally inside Markdown: no links, images, emphasis, colors, or math."""
    text = ' '.join(str(value if value is not None else '').split())
    return _MARKDOWN_SPECIAL.sub(r'\\\1', text)


def can(actor, permission):
    return permission in ROLES.get(actor.role, set())


def token(key):
    return st.session_state.setdefault(f'token:{key}', uuid4().hex)


def mutate(factory, callback, token_key=None):
    try:
        with factory.begin() as session:
            result = callback(session)
            session.flush()
        if token_key:
            st.session_state[f'token:{token_key}'] = uuid4().hex
        st.success('Saved successfully.')
        return result if result is not None else True
    except BusinessError as exc:
        st.error(md_escape(exc))  # Messages can quote user input such as an import reference.
    except IntegrityError:
        st.error('A duplicate or conflicting record was detected. Refresh and check existing records.')
    except SQLAlchemyError as exc:
        log_failure('Database action failed', exc)
        st.error('The database action failed. No partial changes were saved. Check your connection and retry.')
    return None


def options(factory, model, label, predicate=None):
    with factory() as session:
        query = select(model)
        if predicate is not None:
            query = query.where(predicate)
        records = session.scalars(query).all()
        return {r.id: label(r) for r in records}


def choose(label, values, key=None, optional=False):
    data = {None: 'None / All', **values} if optional else values
    if not data:
        st.info(f'Create records first: {label}.')
        return None
    return st.selectbox(label, list(data), format_func=data.get, key=key)


def uploaded(file):
    if file is None:
        return None
    return file.name, file.type, file.getvalue()


def excel_bytes(frame):
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine='openpyxl') as writer:
        frame.to_excel(writer, index=False, sheet_name='Report')
        sheet = writer.sheets['Report']
        from openpyxl.styles import Font, PatternFill
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='18324B')
            cell.font = Font(color='FFFFFF', bold=True)
        from openpyxl.utils import get_column_letter
        for i, column in enumerate(frame.columns, 1):
            sheet.column_dimensions[get_column_letter(i)].width = min(45, max(16, len(str(column)) + 3))
        # Prevent spreadsheet formula injection from names and descriptions.
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    return stream.getvalue()


def safe_csv(frame):
    copy = frame.copy()
    for col in copy.select_dtypes(include=['object', 'string']).columns:
        copy[col] = copy[col].map(lambda value: "'" + value if isinstance(value, str) and value[:1] in ('=', '+', '-', '@', '\t', '\r') else value)
    return copy.to_csv(index=False).encode('utf-8-sig')


def pdf_bytes(frame, title):
    import nova.documents  # Registers the bundled fonts for identical rendering on every machine.
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    buffer = io.BytesIO()
    page = landscape(A3)
    doc = SimpleDocTemplate(buffer, pagesize=page, leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    styles['Heading1'].fontName = 'NovaSansBold'
    text_style = ParagraphStyle('cell', fontName='NovaSans', fontSize=7, leading=9)
    cells = [[Paragraph(html.escape(str(c)), text_style) for c in frame.columns]]
    for row in frame.itertuples(index=False, name=None):
        cells.append([Paragraph(html.escape(f'{x:,.2f}' if isinstance(x, float) else str(x)), text_style) for x in row])
    table = Table(cells, repeatRows=1, splitInRow=1, colWidths=[(page[0] - 48) / max(len(frame.columns), 1)] * len(frame.columns))
    table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#DDEAF6')),
                               ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('GRID', (0, 0), (-1, -1), .2, colors.lightgrey),
                               ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F5F7FA')])]))
    def footer(canvas, document):
        canvas.setFont('NovaSans', 8)
        canvas.drawString(24, 12, f'Nova ERP | {title} | Page {document.page}')
    doc.build([Paragraph(html.escape(title), styles['Heading1']), Spacer(1, 10), table], onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def table(frame, name, key=None, hide=None):
    key = key or name.replace(' ', '_')
    if frame.empty:
        st.info('No records match the selected filters.')
        return
    frame = frame.drop(columns=[c for c in (hide or []) if c in frame.columns]).copy()
    if 'SR No.' not in frame.columns:
        frame.insert(0, 'SR No.', range(1, len(frame) + 1))
    query = st.text_input('Search this table', key=f'search:{key}')
    if query:
        mask = frame.astype(str).apply(lambda col: col.str.contains(query, case=False, regex=False)).any(axis=1)
        frame = frame[mask]
    visible = st.multiselect('Visible columns', list(frame.columns), default=list(frame.columns), key=f'columns:{key}')
    if not visible:
        st.info('Select at least one column.')
        return
    shown = frame[visible]
    signature = hashlib.sha256(safe_csv(shown)).hexdigest()
    if st.session_state.get(f'pdf_signature:{key}') != signature:
        st.session_state.pop(f'pdf:{key}', None)
    st.session_state[f'pdf_signature:{key}'] = signature
    cols = st.columns(4)
    cols[0].download_button('Download CSV', safe_csv(shown), f'{key}.csv', 'text/csv', key=f'csv:{key}')
    cols[1].download_button('Download Excel', excel_bytes(shown), f'{key}.xlsx', key=f'xlsx:{key}')
    # Large exports use CSV/Excel; PDF generation is an explicit user action.
    if len(shown) <= 2000 and cols[2].button('Prepare PDF / Print', key=f'makepdf:{key}'):
        st.session_state[f'pdf:{key}'] = pdf_bytes(shown, name)
    if f'pdf:{key}' in st.session_state:
        cols[3].download_button('Download PDF', st.session_state[f'pdf:{key}'], f'{key}.pdf', 'application/pdf', key=f'pdfdownload:{key}')
    total_pages = max(1, (len(shown) + 99) // 100)
    page = st.number_input(f'Page (100 rows) · {len(shown):,} records', min_value=1, max_value=total_pages, value=1, key=f'page:{key}:{total_pages}')
    displayed = shown.iloc[(page - 1) * 100:page * 100]
    def status_style(row):
        flags = [str(row[c]) for c in row.index if 'status' in c.lower() or c == 'Booking Voucher']
        red = any(v in ('Pending', 'Overdue', 'Rejected', 'Cancelled') for v in flags)
        return ['background-color:#fff0f0;color:#a82127' if red else '' for _ in row]
    st.dataframe(displayed.style.apply(status_style, axis=1), hide_index=True, width='stretch')
    st.caption('CSV and Excel contain all filtered rows. Open the PDF to print. Numeric balances are in PKR.')
